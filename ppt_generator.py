"""요약 결과를 .pptx 파일로 변환."""
from __future__ import annotations

import io
import re
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


BRAND_NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BRAND_BLUE = RGBColor(0x1F, 0x6F, 0xEB)
TEXT_DARK = RGBColor(0x22, 0x22, 0x22)


def _split_summary(summary: str) -> tuple[list[str], list[str]]:
    """요약 텍스트를 (전체 불릿, 핵심 불릿 5개) 로 나눈다.

    Gemma 응답에서 '핵심 불릿' 섹션을 별도로 추출. 없으면 전체에서
    상위 5개를 핵심으로 사용한다.
    """
    lines = [ln.strip() for ln in summary.splitlines()]
    bullets: list[str] = []
    key_bullets: list[str] = []

    in_key_section = False
    for ln in lines:
        if not ln:
            continue
        if re.search(r"핵심\s*불릿", ln):
            in_key_section = True
            continue
        # 마크다운 헤더는 스킵
        if ln.startswith("#"):
            in_key_section = False
            continue
        # 불릿 마커 정규화
        m = re.match(r"^\s*(?:[-*•●▪]|\d+[.)])\s*(.+)", ln)
        text = m.group(1).strip() if m else ln
        # 마크다운 강조 제거
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)

        if in_key_section:
            key_bullets.append(text)
        else:
            bullets.append(text)

    if not key_bullets:
        key_bullets = bullets[:5]
    return bullets, key_bullets[:5]


def _add_title_slide(prs: Presentation, article_title: str, source: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # 배경 띠
    left, top, width, height = Inches(0), Inches(0), prs.slide_width, Inches(1.2)
    bar = slide.shapes.add_shape(1, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_NAVY
    bar.line.fill.background()

    # 헤더 텍스트
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.7))
    p = tx.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "선행기구개발그룹 뉴스 레이더"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 기사 제목
    tx2 = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(12), Inches(2.5))
    tf = tx2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = article_title
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = TEXT_DARK

    # 출처 / 날짜
    tx3 = slide.shapes.add_textbox(Inches(0.7), Inches(5.5), Inches(12), Inches(0.6))
    p = tx3.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = f"출처: {source or 'N/A'}    |    생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    run.font.size = Pt(14)
    run.font.color.rgb = BRAND_BLUE


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 좌측 컬러 바
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.25), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_BLUE
    bar.line.fill.background()

    # 슬라이드 타이틀
    tx = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.9))
    p = tx.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = BRAND_NAVY

    # 본문 박스
    body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(12), Inches(5.8))
    tf = body_box.text_frame
    tf.word_wrap = True

    if not bullets:
        bullets = ["(요약 내용 없음)"]

    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = f"• {b}"
        run.font.size = Pt(16)
        run.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)


def _add_closing_slide(prs: Presentation, link: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tx = slide.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(12), Inches(2))
    tf = tx.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "원문 링크"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = BRAND_NAVY

    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = link or "(링크 없음)"
    run2.font.size = Pt(14)
    run2.font.color.rgb = BRAND_BLUE


def build_pptx(
    article_title: str,
    article_link: str,
    article_source: str,
    summary_markdown: str,
) -> bytes:
    """요약 마크다운을 받아 .pptx 바이너리를 반환."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    bullets, key_bullets = _split_summary(summary_markdown)

    _add_title_slide(prs, article_title, article_source)
    _add_bullet_slide(prs, "Executive Summary - 핵심 5개", key_bullets)

    # 전체 불릿을 페이지당 ~7개씩 분할
    page_size = 7
    for i in range(0, len(bullets), page_size):
        chunk = bullets[i : i + page_size]
        page_idx = i // page_size + 1
        _add_bullet_slide(prs, f"기술 분석 ({page_idx})", chunk)

    _add_closing_slide(prs, article_link)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
